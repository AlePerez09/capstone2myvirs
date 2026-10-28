"""
Tests for word_lists.py using mock Excel files built with the layout
Isa described for Eric's lists. Run from the repo folder:

    python tests/test_word_lists.py

Replace with a check against the real files once we have them.
"""

import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import word_lists  # noqa: E402

YELLOW = PatternFill("solid", fgColor="FFFFFF00")
BLUE = PatternFill("solid", fgColor="FF5B9BD5")
GREEN = PatternFill("solid", fgColor="FF70AD47")


def band_file(path, rows, total=None):
    wb = Workbook()
    ws = wb.active
    ws.append(["Frequency", "Word"])
    for freq, word in rows:
        ws.append([freq, word])
    if total is not None:
        ws.append([total, None])  # total row with no word
    wb.save(path)


def grade_grid(path, columns, headers=None):
    """columns: {grade header: [(word, fill or None), ...]}"""
    wb = Workbook()
    ws = wb.active
    headers = headers or list(columns)
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
        for r, (word, fill) in enumerate(columns[h], 2):
            cell = ws.cell(row=r, column=c, value=word)
            if fill:
                cell.fill = fill
    wb.save(path)


def make_mock(folder):
    folder = Path(folder)
    band_file(folder / "New-High-Frequency-3450.xlsx",
              [(5000, "the"), (900, "Cell"), (700, "water"), (300, "analyze")])
    band_file(folder / "MID-5005.xlsx",
              [(50, "photosynthesis"), (40, "glucose"), (30, "water")], total=5005)
    band_file(folder / "LOW-8080.xlsx", [(5, "mitochondria"), (2, "in spite of")], total=8080)
    grade_grid(folder / "K-12-significant-words.xlsx", {
        "K": [("water", None), ("", None)],
        "1": [("cell", None)],
        "4": [("analyze", YELLOW)],
        "7": [("photosynthesis", GREEN), ("equation", BLUE)],
    })
    grade_grid(folder / "K-12-math-science-words.xlsx", {
        "3": [("cell", GREEN)],
        "6": [("equation", BLUE), ("data", YELLOW)],
        "9": [("photosynthesis", GREEN)],
    })


def test_lexicon():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        make_mock(tmp)
        assert word_lists.files_present(tmp)
        lex = word_lists.load_lexicon(tmp)

        assert lex["the"]["band"] == "HIGH" and lex["the"]["frequency_count"] == 5000
        assert lex["cell"]["band"] == "HIGH", "words are lowercased"
        assert lex["water"]["band"] == "HIGH", "a word in two bands keeps the higher one"
        assert lex["photosynthesis"]["band"] == "MID"
        assert lex["mitochondria"]["band"] == "LOW"
        assert "in spite of" in lex, "multiword entries are kept"
        assert not any(w.isdigit() for w in lex), "total rows are skipped"

        assert lex["water"]["grade_level"] == "K" and lex["water"]["subject"] == "general"
        assert lex["analyze"]["subject"] == "academic", "yellow = academic in the significant list"
        assert lex["cell"]["grade_level"] == "G1", "earliest grade wins"
        assert lex["cell"]["subject"] == "science", "math-science file is more specific"
        assert lex["photosynthesis"]["grade_level"] == "G7"
        assert lex["equation"]["subject"] == "math"
        assert lex["data"]["subject"] == "math_science", "yellow = both in the math-science list"
        assert lex["data"]["band"] == "NONE", "grade-list words not in a band file get NONE"


def test_unknown_color_fails():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        make_mock(tmp)
        red = PatternFill("solid", fgColor="FFFF0000")
        grade_grid(Path(tmp) / "K-12-significant-words.xlsx", {"K": [("water", red)],
                                                                "1": [], "2": []})
        try:
            word_lists.load_lexicon(tmp)
        except word_lists.WordListError as e:
            assert "unrecognized" in str(e)
        else:
            raise AssertionError("an unknown color should raise an error")


def test_grade_headers():
    for header, expected in [("K", "K"), ("Kindergarten", "K"), ("1", "G1"),
                             (12, "G12"), ("Grade 5", "G5"), ("G10", "G10"), ("13", None)]:
        assert word_lists.parse_grade(header) == expected, header


def test_ranking_uses_lexicon():
    import priority_ranking

    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        make_mock(tmp)
        lex = word_lists.load_lexicon(tmp)
        ranked = priority_ranking.rank_text(
            "Photosynthesis happens in the cell. The cell uses water and glucose.",
            lexicon=lex, text_grade="G7")
        bands = {r["headword"]: r["band"] for r in ranked}
        assert bands["photosynthesis"] == "MID" and bands["cell"] == "HIGH"
        assert bands["happen"] == "NONE"


def test_virs_csv():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        path = Path(tmp) / "virs_words.csv"
        path.write_text(
            "value,category,grade\n"
            "cell,k1,G4\ncell,stem,G4\ncell,hi,G4\n"
            "the,k1,\n"
            "analyze,k1,G2\nanalyze,awl,G2\nanalyze,hi,G2\n"
            "glucose,k2,G6\n"
            "habitat,med,G3\nhabitat,med,G5\n"
            "zygote,low,\n", encoding="utf-8")
        lex = word_lists.load_virs_csv(path)
        assert lex["cell"] == {"band": "HIGH", "subject": "stem", "k_level": "K1",
                               "grade_level": "G4", "frequency_count": None}
        assert lex["the"]["band"] == "HIGH", "K1-only words count as HIGH"
        assert lex["analyze"]["subject"] == "academic"
        assert lex["glucose"]["band"] == "HIGH" and lex["glucose"]["k_level"] == "K2"
        assert lex["habitat"]["band"] == "MID" and lex["habitat"]["grade_level"] == "G3", "earliest grade"
        assert lex["zygote"]["band"] == "LOW" and lex["zygote"]["grade_level"] is None


if __name__ == "__main__":
    tests = [v for k, v in dict(globals()).items() if k.startswith("test_")]
    for t in tests:
        t()
        print(f"passed  {t.__name__}")
    print(f"\nAll {len(tests)} tests passed")
