"""
MyVIRS Team A, Sprint 3: load Eric's word lists into one lookup table.

Reads the five Excel files Eric shared in the Project channel and builds a
"lexicon": {headword: {"band", "subject", "grade_level", "frequency_count"}},
the format priority_ranking.py and contract.py expect.

Put the files in data/word_lists/:
    New-High-Frequency-3450.xlsx   HIGH band  (columns: Frequency, Word)
    MID-5005.xlsx                  MID band
    LOW-8080.xlsx                  LOW band
    K-12-significant-words.xlsx    grade grid, subject by cell color
    K-12-math-science-words.xlsx   grade grid, subject by cell color

Layout, as described in Isa's integration layer design:
- Band files have a Frequency column and a Word column. MID and LOW end
  with a total row that has no word; rows without a word are skipped.
- K-12 files are grids with one column per grade (K, then grades 1-12)
  and many blank cells. Subject is stored only as the cell's fill color:
    significant words:   no fill = general, yellow = academic (PSAWL),
                         blue = math, green = science
    math-science words:  blue = math, green = science, yellow = both
  Yellow means different things in the two files, so each file has its
  own color map. An unrecognized color stops with an error instead of
  guessing.
- A word's grade_level is the earliest grade it appears in, since grade
  lists build on the ones before them.

NOTE: written before we had the real files, from Isa's description, and
tested on mock files with that layout. If the real files differ (header
names, colors, sheet layout), adjust COLUMN_NAMES, COLOR_MAPS or
classify_color() below.

Run:
    python3 word_lists.py                 # summary of what loaded
    python3 word_lists.py --word cell     # look up one word
"""

import argparse
import re
import sys
from pathlib import Path

WORD_LIST_DIR = Path(__file__).parent / "data" / "word_lists"

# Filename patterns, so slightly different names still match
BAND_FILES = {"HIGH": "*high*.xlsx", "MID": "*mid*.xlsx", "LOW": "*low*.xlsx"}
SIGNIFICANT_FILE = "*significant*.xlsx"
MATH_SCIENCE_FILE = "*math*science*.xlsx"

COLUMN_NAMES = {"word": ("word", "words", "headword"),
                "frequency": ("frequency", "freq", "count")}

GRADES = ["K"] + [f"G{i}" for i in range(1, 13)]
BAND_ORDER = ["HIGH", "MID", "LOW"]

# color name -> subject, per file
COLOR_MAPS = {
    "significant": {None: "general", "yellow": "academic", "blue": "math", "green": "science"},
    "math_science": {"blue": "math", "green": "science", "yellow": "math_science"},
}


class WordListError(Exception):
    pass


# ---------- Helpers ----------

def find_file(folder, pattern):
    """Case-insensitive filename match (skips Excel's ~$ lock files)."""
    import fnmatch

    folder = Path(folder)
    if not folder.is_dir():
        return None
    matches = sorted(p for p in folder.iterdir()
                     if fnmatch.fnmatch(p.name.lower(), pattern.lower())
                     and not p.name.startswith("~$"))
    return matches[0] if matches else None


def clean_word(value):
    if value is None:
        return None
    word = re.sub(r"\s+", " ", str(value)).strip().lower()
    return word or None


def parse_grade(header):
    """'K', 'Kindergarten', '1', 'G1', 'Grade 1', 1 -> 'K' or 'G1'."""
    if header is None:
        return None
    text = str(header).strip().lower()
    if text in ("k", "kg", "kindergarten") or text.startswith("kinder"):
        return "K"
    m = re.fullmatch(r"(?:g|gr|grade)?\s*\.?\s*(\d{1,2})(?:\.0)?", text)
    if m and 1 <= int(m.group(1)) <= 12:
        return f"G{int(m.group(1))}"
    return None


def classify_color(cell):
    """Return None (no fill), 'yellow', 'blue', 'green', or raise."""
    fill = cell.fill
    if fill is None or fill.fill_type in (None, "none"):
        return None
    color = fill.fgColor
    if color is None:
        return None
    if color.type == "rgb" and color.rgb:
        rgb = str(color.rgb)[-6:]
        if rgb.upper() in ("000000", "FFFFFF"):  # default / white fill
            return None
        r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
        if r > 180 and g > 180 and b < 150:
            return "yellow"
        if b > r + 30 and b > g - 10:
            return "blue"
        if g > r + 30 and g > b + 10:
            return "green"
        raise WordListError(f"cell {cell.coordinate}: unrecognized fill color #{rgb}")
    if color.type == "indexed" and color.indexed in (64,):  # system default
        return None
    raise WordListError(
        f"cell {cell.coordinate}: fill color is a {color.type} color "
        f"({color.theme if color.type == 'theme' else color.indexed}), not RGB. "
        "Add it to classify_color()."
    )


# ---------- Band files ----------

def load_band_file(path):
    """Return {word: frequency} from a Frequency/Word sheet."""
    from openpyxl import load_workbook

    # Read-only workbooks keep the file open until closed, and Windows
    # won't let anyone delete or move an open file, so always close it
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()
    word_col = freq_col = None
    header_row = 0
    for i, row in enumerate(rows[:10]):
        names = [str(v).strip().lower() if v is not None else "" for v in row]
        for j, name in enumerate(names):
            if name in COLUMN_NAMES["word"]:
                word_col = j
            elif name in COLUMN_NAMES["frequency"]:
                freq_col = j
        if word_col is not None:
            header_row = i + 1
            break
    if word_col is None:
        raise WordListError(f"{path.name}: no 'Word' column header found")

    words = {}
    for row in rows[header_row:]:
        word = clean_word(row[word_col] if word_col < len(row) else None)
        if not word or word.replace(".", "").isdigit():
            continue  # blank or total row
        freq = row[freq_col] if freq_col is not None and freq_col < len(row) else None
        try:
            freq = int(freq) if freq is not None else None
        except (TypeError, ValueError):
            freq = None
        words.setdefault(word, freq)
    return words


# ---------- K-12 grade grids ----------

def load_grade_grid(path, color_map_name):
    """Return {word: (grade, subject)} using each word's earliest grade."""
    from openpyxl import load_workbook

    color_map = COLOR_MAPS[color_map_name]
    wb = load_workbook(path, data_only=True)  # not read-only: we need the fills
    try:
        return _read_grade_grid(path, wb.worksheets[0], color_map)
    finally:
        wb.close()


def _read_grade_grid(path, ws, color_map):
    grade_cols, header_row = {}, None
    for row in ws.iter_rows(min_row=1, max_row=10):
        found = {c.column: parse_grade(c.value) for c in row if parse_grade(c.value)}
        if len(found) >= 3:
            grade_cols, header_row = found, row[0].row
            break
    if not grade_cols:
        raise WordListError(f"{path.name}: no grade header row (K, 1, 2, ...) found")

    words = {}
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            grade = grade_cols.get(cell.column)
            word = clean_word(cell.value)
            if not grade or not word:
                continue
            color = classify_color(cell)
            if color not in color_map:
                raise WordListError(f"{path.name} {cell.coordinate} ('{word}'): color "
                                    f"'{color}' has no meaning in this file")
            subject = color_map[color]
            old = words.get(word)
            if old is None or GRADES.index(grade) < GRADES.index(old[0]):
                words[word] = (grade, subject)
    return words


# ---------- Lexicon ----------

def files_present(folder=WORD_LIST_DIR):
    return all(find_file(folder, p) for p in
               [*BAND_FILES.values(), SIGNIFICANT_FILE, MATH_SCIENCE_FILE])


def load_lexicon(folder=WORD_LIST_DIR):
    folder = Path(folder)
    lexicon = {}

    for band in BAND_ORDER:
        path = find_file(folder, BAND_FILES[band])
        if path is None:
            raise WordListError(f"missing {band} file ({BAND_FILES[band]}) in {folder}")
        for word, freq in load_band_file(path).items():
            if word not in lexicon:  # a word in two band files keeps the higher band
                lexicon[word] = {"band": band, "subject": "general",
                                 "grade_level": None, "frequency_count": freq}

    for pattern, map_name in [(SIGNIFICANT_FILE, "significant"),
                              (MATH_SCIENCE_FILE, "math_science")]:
        path = find_file(folder, pattern)
        if path is None:
            raise WordListError(f"missing {pattern} in {folder}")
        for word, (grade, subject) in load_grade_grid(path, map_name).items():
            entry = lexicon.setdefault(word, {"band": "NONE", "subject": "general",
                                              "grade_level": None, "frequency_count": None})
            if entry["grade_level"] is None or GRADES.index(grade) < GRADES.index(entry["grade_level"]):
                entry["grade_level"] = grade
            # math/science file is more specific than the general list
            if map_name == "math_science" or entry["subject"] == "general":
                entry["subject"] = subject
    return lexicon


def main():
    parser = argparse.ArgumentParser(description="Load Eric's word lists")
    parser.add_argument("--folder", default=str(WORD_LIST_DIR))
    parser.add_argument("--word", help="look up one word")
    args = parser.parse_args()

    if not files_present(args.folder):
        print(f"Eric's five word list files aren't in {args.folder} yet.")
        sys.exit(1)

    lexicon = load_lexicon(args.folder)
    if args.word:
        print(args.word, lexicon.get(args.word.lower(), "not in any list"))
        return

    from collections import Counter
    print(f"Loaded {len(lexicon):,} words")
    print("Bands:   ", dict(Counter(e["band"] for e in lexicon.values())))
    print("Subjects:", dict(Counter(e["subject"] for e in lexicon.values())))
    print("With a grade:", sum(1 for e in lexicon.values() if e["grade_level"]))


if __name__ == "__main__":
    main()
